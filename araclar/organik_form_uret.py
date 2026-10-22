import sys,os,copy
import openpyxl
from openpyxl.styles import PatternFill,Font,Alignment,Border,Side
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor,AnchorMarker
from openpyxl.utils import get_column_letter as CL
import content as C
G='006600';LG='EBF1DE'
GF=PatternFill('solid',fgColor=G);LF=PatternFill('solid',fgColor=LG)
thin=Side(style='thin');med=Side(style='medium')
A='/tmp/claude-0/-home-user/5d050a5c-9efe-5868-b7c0-331c83134422/scratchpad/w/assets/'
def anchor(c0,o0,r0,ro0,c1,o1,r1,ro1):
    a=TwoCellAnchor(editAs='oneCell');a._from=AnchorMarker(col=c0,colOff=o0,row=r0,rowOff=ro0);a.to=AnchorMarker(col=c1,colOff=o1,row=r1,rowOff=ro1);return a
def spans(weights,total=20):
    s=[max(1,round(w*total/sum(weights))) for w in weights]
    while sum(s)>total:
        k=s.index(max(s));s[k]-=1
    while sum(s)<total:
        k=len(s)-1 if len(s)>1 else 0;s[k]+=1
    return s
def box(ws,r1,c1,r2,c2,val=None,fill=None,font=None,al=None,outer=True):
    if (r2>r1) or (c2>c1): ws.merge_cells(start_row=r1,start_column=c1,end_row=r2,end_column=c2)
    for r in range(r1,r2+1):
        for c in range(c1,c2+1):
            cell=ws.cell(r,c)
            cell.border=Border(left=med if c==1 else thin if c==c1 else None,right=thin if c==c2 else None,top=thin if r==r1 else None,bottom=thin if r==r2 else None)
            if fill: cell.fill=fill
    cell=ws.cell(r1,c1)
    if val is not None: cell.value=val
    if font: cell.font=font
    if al: cell.alignment=al
def build(code,tr,en,ar,info,cols,nrows,checks,sigs,out,notes=True,title_note=None):
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Form'
    for c in range(1,21): ws.column_dimensions[CL(c)].width=9
    ws.column_dimensions['R'].width=10.2;ws.column_dimensions['S'].width=10.66;ws.column_dimensions['T'].width=17.2
    ws.sheet_format.defaultRowHeight=14.4
    # header
    for r in range(1,5):
        ws.row_dimensions[r].height=24
        for c in range(1,21): ws.cell(r,c).fill=GF
    ln=max(len(en),1)
    sz=28 if ln<=30 else 24 if ln<=36 else 20 if ln<=44 else 17
    box(ws,1,4,4,15,'%s\n%s'%(en.upper(),ar),GF,Font(name='Calibri',sz=sz,b=True,color='FFFFFF'),Alignment(horizontal='center',vertical='center',wrap_text=True))
    meta=[('DATE / التاريخ','….../….../.…'),('FORM NUMBER / رقم النموذج',code),('REVISION / المراجعة','R00'),('PUBLIC. DATE / تاريخ النشر','22.10.2026')]
    for k,(l,v) in enumerate(meta,1):
        box(ws,k,18,k,19,l,GF,Font(name='Calibri',sz=9,b=True,color='FFFFFF'),Alignment(horizontal='left',vertical='center',wrap_text=True))
        box(ws,k,20,k,20,v,GF,Font(name='Calibri',sz=10,b=(k==2),color='FFFFFF'),Alignment(horizontal='left',vertical='center',wrap_text=True))
    for r in range(1,5):
        ws.cell(r,1).border=Border(left=med,top=med if r==1 else None)
    def bar(r,t):
        ws.row_dimensions[r].height=20
        box(ws,r,1,r,20,t,GF,Font(name='Calibri',sz=11,b=True,color='FFFFFF'),Alignment(horizontal='left',vertical='center',wrap_text=True))
    r=6;bar(r,'GENERAL INFORMATION / المعلومات العامة');r+=1
    lab=Font(name='Calibri',sz=11,b=True,color='000000');lal=Alignment(horizontal='left',vertical='top',wrap_text=True)
    for k in range(0,len(info),4):
        grp=info[k:k+4];ws.row_dimensions[r].height=34.05
        sp=[5]*len(grp)
        if len(grp)<4: sp[-1]+=5*(4-len(grp))
        c=1
        for (e,a),s in zip(grp,sp):
            box(ws,r,c,r,c+s-1,'%s\n%s'%(e,a),LF,lab,lal);c+=s
        r+=1
    r+=1
    bar(r,cols['title']);r+=1
    cl=cols['cols'];sp=[2]+spans([w for _,_,w in cl],18)
    ws.row_dimensions[r].height=37.2
    hf=Font(name='Calibri',sz=11,b=True,color='000000');hal=Alignment(horizontal='center',vertical='center',wrap_text=True)
    c=1;box(ws,r,1,r,2,'No',LF,hf,hal);c=3
    for (e,a,w),s in zip(cl,sp[1:]):
        box(ws,r,c,r,c+s-1,'%s\n%s'%(e,a),LF,hf,hal);c+=s
    r+=1
    df=Font(name='Calibri',sz=10,color='000000');dal=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for n in range(1,nrows+1):
        ws.row_dimensions[r].height=22.05
        box(ws,r,1,r,2,n,None,df,dal);c=3
        for s in sp[1:]:
            box(ws,r,c,r,c+s-1,None,None,df,dal);c+=s
        r+=1
    r+=1
    if checks:
        bar(r,checks['title']);r+=1
        ws.row_dimensions[r].height=30
        box(ws,r,1,r,20,'   '.join('☐ %s / %s'%(e,a) for e,a in checks['items']),None,Font(name='Calibri',sz=10,color='000000'),Alignment(horizontal='left',vertical='center',wrap_text=True));r+=2
    if notes:
        bar(r,'NOTES / ملاحظات');r+=1
        ws.row_dimensions[r].height=16;ws.row_dimensions[r+1].height=16
        box(ws,r,1,r+1,20,'Gaps found and corrective actions.',None,Font(name='Calibri',sz=9,color='000000'),Alignment(horizontal='left',vertical='center',wrap_text=True));r+=3
    # imza
    n=len(sigs);w=[6,6,8] if n==3 else [10,10]
    c=1;hr=r
    for (e,a),ww in zip(sigs,w):
        box(ws,hr,c,hr,c+ww-1,'%s / %s'%(e,a),GF,Font(name='Calibri',sz=11,b=True,color='FFFFFF'),Alignment(horizontal='center',vertical='center',wrap_text=True))
        box(ws,hr+1,c,hr+3,c+ww-1,None,None,None,None);c+=ww
    ws.row_dimensions[hr].height=20
    r=hr+4
    ws.row_dimensions[r].height=30
    box(ws,r,1,r,20,'Retention: at least 5 years or per certification body / customer requirement. Any nonconformity opens a corrective action (DÖF). | الاحتفاظ: 5 سنوات على الأقل أو حسب متطلبات جهة الاعتماد / العميل. وتفتح أي حالة عدم مطابقة إجراءً تصحيحيًا (DÖF).',GF,Font(name='Calibri',sz=11,b=True,color='FFFFFF'),Alignment(horizontal='left',vertical='center',wrap_text=True))
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=9;ws.sheet_properties.pageSetUpPr=openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True);ws.page_setup.fitToHeight=0
    ws.page_margins.left=ws.page_margins.right=0.4;ws.page_margins.top=ws.page_margins.bottom=0.5
    im=Image(A+'ems_white.png');im.anchor=anchor(0,238124,0,198193,2,171449,3,205802);ws.add_image(im)
    im=Image(A+'gots.png');im.anchor=anchor(15,66676,0,28575,16,609600,3,283176);ws.add_image(im)
    wb.save(out)
