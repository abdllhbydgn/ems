#!/usr/bin/env python3
"""Organik prosedür/talimat/politika (EMS.OR.*) Word başlığına GOTS logosunu EMS logosunun yanına ekler (Seçenek B, kullanıcı onayı 27.10.2026).
Logo hücresi 450 twip genişler, başlık hücresi o kadar daralır; EMS 760000 EMU, GOTS 700000 EMU. İdempotent (hücrede 2 görsel varsa dokunmaz).
Kullanım: python3 araclar/organik_docx_gots.py"""
import json, hashlib, docx, os
from docx.shared import Emu, Twips
from docx.oxml.ns import qn
ROOT = "prosedurler"; PNG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "gots_hd_kare.png")
def fix(path):
    d = docx.Document(path); t = d.sections[0].header.tables[0]; row = t.rows[0]; c = row.cells[0]
    if sum(1 for e in c._tc.iter() if e.tag.endswith("}extent")) >= 2: return False
    for e in c._tc.iter():
        tg = e.tag.split("}")[1]
        if tg == "extent":
            cx, cy = int(e.get("cx")), int(e.get("cy")); e.set("cx", "760000"); e.set("cy", str(int(cy * 760000 / cx)))
        if tg == "ext" and e.get("cx"):
            cx, cy = int(e.get("cx")), int(e.get("cy")); k = 760000 / 723900; e.set("cx", str(int(cx * k))); e.set("cy", str(int(cy * k)))
    r = c.paragraphs[0].add_run(" "); r.add_picture(PNG, width=Emu(700000))
    tw = lambda emu: int(emu / 635)
    w0, w1 = tw(int(row.cells[0].width)), tw(int(row.cells[1].width))
    row.cells[0].width = Twips(w0 + 450); row.cells[1].width = Twips(w1 - 450)
    g = t._tbl.tblGrid.findall(qn("w:gridCol")); g[0].set(qn("w:w"), str(w0 + 450)); g[1].set(qn("w:w"), str(w1 - 450))
    d.save(path); return True
if __name__ == "__main__":
    m = json.load(open(ROOT + "/manifest.json")); it = m if isinstance(m, list) else m.get("items", m); n = 0
    for x in it:
        if ".OR." not in x.get("code", ""): continue
        p = os.path.join(ROOT, x["file"])
        if fix(p):
            x["rev"] = "gots-" + hashlib.md5(open(p, "rb").read()).hexdigest()[:8]; n += 1
    json.dump(m, open(ROOT + "/manifest.json", "w"), ensure_ascii=False, indent=1)
    print("güncellenen belge:", n)
